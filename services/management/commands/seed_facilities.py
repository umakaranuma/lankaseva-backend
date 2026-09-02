"""Seeds the full government-facility directory from the curated spreadsheet.

Source: services/management/commands/data/facilities_v6.xlsx
    31 category sheets, one row per office, each with real GPS coordinates.

Every row becomes a Service (+ one primary phone + weekly opening hours).
Idempotent — keyed on the sheet's "Record ID" (stored lower-cased as
`Service.code`), so re-running updates in place.

    python manage.py seed_facilities            # upsert every row
    python manage.py seed_facilities --fresh     # wipe ALL services first
    python manage.py seed_facilities --file other.xlsx

Requires `openpyxl` (in requirements.txt).
"""

import hashlib
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from services.models import (
    Category,
    OpeningHourSlot,
    OpeningHours,
    Service,
    ServicePhone,
)

DATA_DIR = Path(__file__).resolve().parent / "data"
DEFAULT_FILE = DATA_DIR / "facilities_v6.xlsx"

# ── Category metadata: code -> (English name, icon, ARGB colour). ──
# Keeps this command self-sufficient — it ensures every referenced category
# exists (names are only applied on CREATE, so a nicer existing name wins).
_GREEN = "0xFF0D6552"
CATEGORY_STYLE = {
    "secretariat": ("Government Secretariats", "account_balance", _GREEN),
    "transport": ("Transport", "directions_bus", _GREEN),
    "hospital": ("Hospitals", "local_hospital", _GREEN),
    "heritage": ("Museums & Heritage", "museum", _GREEN),
    "environment": ("Environment & Wildlife", "park", _GREEN),
    "employment": ("Employment & Labour", "work", _GREEN),
    "identity": ("Identity & Migration", "badge", _GREEN),
    "tax_revenue": ("Tax & Revenue", "account_balance", _GREEN),
    "agriculture": ("Agriculture & Farming", "agriculture", _GREEN),
    "library": ("Libraries", "local_library", _GREEN),
    "ceb": ("Electricity", "electric_bolt", _GREEN),
    "court": ("Courts", "gavel", _GREEN),
    "fire_rescue": ("Fire & Rescue", "local_fire_department", _GREEN),
    "police": ("Police Stations", "local_police", _GREEN),
    "post": ("Post Offices", "local_post_office", _GREEN),
    "school": ("Schools", "school", _GREEN),
    "water": ("Water Supply", "water_drop", _GREEN),
    "law_enforcement": ("Law Enforcement", "local_police", _GREEN),
    "local_govt": ("Local Government", "location_city", _GREEN),
    "registry": ("Registry & Records", "assignment", _GREEN),
    "standards": ("Standards & Consumer", "verified", _GREEN),
    "emergency": ("Emergency Services", "emergency", "0xFFC62828"),
    "aviation": ("Aviation", "flight", _GREEN),
    "culture": ("Cultural Affairs", "theater_comedy", _GREEN),
    "customs": ("Customs & Border", "local_shipping", _GREEN),
    "energy": ("Energy & Fuel", "local_gas_station", _GREEN),
    "infrastructure": ("Infrastructure & Roads", "construction", _GREEN),
    "land_survey": ("Land & Survey", "map", _GREEN),
    "maritime": ("Maritime & Ports", "directions_boat", _GREEN),
    "mining": ("Mining & Geology", "diamond", _GREEN),
    "social": ("Social Services", "family_restroom", _GREEN),
}

# 24/7 categories — no weekday slots, always-open flag.
ALWAYS_OPEN = {"hospital", "police", "fire_rescue", "emergency", "law_enforcement"}
# is_emergency flag on the service record.
EMERGENCY_FLAG = {"hospital", "police", "fire_rescue", "emergency", "law_enforcement"}
# National hotline added as a secondary number for these categories.
HOTLINE = {
    "police": "119",
    "law_enforcement": "119",
    "fire_rescue": "110",
    "emergency": "117",
    "hospital": "1990",
}

OFFICE_HOURS = [(d, "08:30", "16:15") for d in range(1, 6)]           # Mon–Fri
DEPOT_HOURS = [(d, "05:00", "21:00") for d in range(1, 7)] + [(7, "06:00", "20:00")]

# Sri Lanka bounding box — reject anything outside it.
SL_BBOX = (5.7, 10.0, 79.4, 82.1)  # lat_min, lat_max, lng_min, lng_max


class Command(BaseCommand):
    help = "Seed all government facilities from the curated GPS spreadsheet."

    def add_arguments(self, parser):
        parser.add_argument("--fresh", action="store_true",
                            help="Delete every existing Service before loading.")
        parser.add_argument("--file", default=str(DEFAULT_FILE),
                            help="Path to the .xlsx directory file.")

    @transaction.atomic
    def handle(self, *args, **opts):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("openpyxl is required — `pip install openpyxl`.")

        path = Path(opts["file"])
        if not path.exists():
            raise CommandError(f"Data file not found: {path}")

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows = list(self._iter_rows(wb))
        if not rows:
            raise CommandError("No valid rows found in the spreadsheet.")

        self._ensure_categories(rows)

        if opts["fresh"]:
            deleted = Service.objects.all().delete()[0]
            self.stdout.write(self.style.WARNING(
                f"Cleared {deleted} existing service rows."))

        created = updated = skipped = 0
        for r in rows:
            ok = self._save(r)
            if ok is None:
                skipped += 1
            elif ok:
                created += 1
            else:
                updated += 1

        self.stdout.write(self.style.SUCCESS(
            f"Facilities seeded: {created} created, {updated} updated, "
            f"{skipped} skipped ({Service.objects.count()} total)."))

    # ── spreadsheet → normalised rows ────────────────────────────────

    def _iter_rows(self, wb):
        for sheet in wb.sheetnames:
            if sheet.lower().startswith("taxonomy"):
                continue
            ws = wb[sheet]
            grid = list(ws.iter_rows(values_only=True))
            header_idx = next(
                (i for i, row in enumerate(grid)
                 if row and str(row[0]).strip() == "Record ID"),
                None,
            )
            if header_idx is None:
                continue
            for row in grid[header_idx + 1:]:
                if not row or row[0] is None:
                    continue
                rid = str(row[0]).strip()
                code = (str(row[1]).strip() if row[1] else "").lower()
                district = str(row[3]).strip() if row[3] else ""
                ftype = str(row[4]).strip() if row[4] else ""
                name = str(row[5]).strip() if row[5] else ""
                desc = str(row[8]).strip() if len(row) > 8 and row[8] else ""
                try:
                    lat = float(row[6])
                    lng = float(row[7])
                except (TypeError, ValueError):
                    continue
                if not (SL_BBOX[0] <= lat <= SL_BBOX[1]
                        and SL_BBOX[2] <= lng <= SL_BBOX[3]):
                    continue
                if not (rid and code and district and name):
                    continue
                yield {
                    "code": rid.lower(),
                    "category": code,
                    "district": district,
                    "ftype": ftype,
                    "name": name,
                    "desc": desc,
                    "lat": lat,
                    "lng": lng,
                }

    # ── categories ──────────────────────────────────────────────────

    def _ensure_categories(self, rows):
        seen = {r["category"] for r in rows}
        for code in sorted(seen):
            name, icon, color = CATEGORY_STYLE.get(
                code, (code.replace("_", " ").title(), "account_balance", _GREEN))
            Category.objects.update_or_create(
                code=code,
                defaults=dict(icon=icon, color=color),
                create_defaults=dict(
                    name_en=name, name_si=name, name_ta=name,
                    icon=icon, color=color,
                ),
            )
        self.stdout.write(f"Ensured {len(seen)} categories.")

    # ── one service ────────────────────────────────────────────────

    def _save(self, r):
        cat = r["category"]
        address = r["desc"] or f"{r['name']}, {r['district']}, Sri Lanka"
        dept = r["ftype"] or r["category"].replace("_", " ").title()

        _, was_created = Service.objects.update_or_create(
            code=r["code"],
            defaults=dict(
                name_en=r["name"], name_si=r["name"], name_ta=r["name"],
                department_en=dept, department_si=dept, department_ta=dept,
                category_id=cat,
                district=r["district"],
                address_en=address, address_si=address, address_ta=address,
                lat=r["lat"], lng=r["lng"],
                website=None, whatsapp=None,
                is_emergency=cat in EMERGENCY_FLAG,
            ),
        )
        service = Service.objects.get(code=r["code"])
        self._save_phones(service, cat, r["code"])
        self._save_hours(service, cat, r["ftype"])
        return was_created

    def _save_phones(self, service, cat, code):
        service.phones.all().delete()
        # Deterministic placeholder landline (stable across re-seeds).
        digest = int(hashlib.sha1(code.encode()).hexdigest(), 16)
        area = 11 + digest % 78
        num = f"0{area:02d}{2000000 + digest % 7999999}"
        phones = [ServicePhone(
            service=service, number=num, is_primary=True,
            label_en="General", label_si="ප්‍රධාන", label_ta="பொது")]
        if cat in HOTLINE:
            phones.append(ServicePhone(
                service=service, number=HOTLINE[cat], is_primary=False,
                label_en="Emergency", label_si="හදිසි", label_ta="அவசரம்"))
        ServicePhone.objects.bulk_create(phones)

    def _save_hours(self, service, cat, ftype):
        record, _ = OpeningHours.objects.update_or_create(
            service=service,
            defaults=dict(is_always_open=cat in ALWAYS_OPEN, notes=None),
        )
        record.slots.all().delete()
        if cat in ALWAYS_OPEN:
            return
        is_depot = cat == "transport" and any(
            k in ftype.lower() for k in ("depot", "stand", "terminal", "station"))
        table = DEPOT_HOURS if is_depot else OFFICE_HOURS
        OpeningHourSlot.objects.bulk_create([
            OpeningHourSlot(hours=record, weekday=wd,
                            open_time=o, close_time=c)
            for wd, o, c in table
        ])
